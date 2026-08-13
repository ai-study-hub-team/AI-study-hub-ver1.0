from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("docs/Bang_tom_tat_phan_tich_bao_mat.docx")
FONT = "Arial"
NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "DCE6F1"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
RED = "9C0006"
AMBER = "9C6500"
GREEN = "006100"


FINDINGS = [
    ("1", "Không giới hạn số lần nhập sai mật khẩu", "Backend",
     "API POST /api/auth/login chỉ kiểm tra mật khẩu rồi trả 401. Không có bộ đếm số lần sai, khóa tạm thời, giới hạn theo IP hoặc CAPTCHA.",
     "Trung bình - cao", "Kẻ xấu có thể thử mật khẩu liên tục (brute force)."),
    ("2", "Có trạng thái LOCKED nhưng không tự động sử dụng", "Backend",
     "UserStatus có LOCKED, nhưng đăng nhập sai không làm tài khoản chuyển sang trạng thái này.",
     "Trung bình", "Chức năng khóa tài khoản chưa hoàn chỉnh."),
    ("3", "Thiếu dữ liệu quản lý đăng nhập thất bại", "Backend / Database",
     "Bảng users không có failed_login_attempts, last_failed_login_at hoặc locked_until.",
     "Trung bình", "Không đủ dữ liệu để triển khai khóa tạm theo tài khoản."),
    ("4", "Không có rate limit cho API đăng nhập", "Backend / Hạ tầng",
     "/api/auth/login là API public và không giới hạn số request theo IP hoặc email trong một khoảng thời gian.",
     "Cao khi triển khai public", "Có nguy cơ brute force và làm tốn CPU do BCrypt."),
    ("5", "Route /app/** không có guard đăng nhập chung", "Frontend",
     "Các route dashboard, upload, library, chat, summary và quiz render DashboardLayout trực tiếp, không được bọc bởi RequireAuth/ProtectedRoute.",
     "Trung bình", "Copy URL sang trình duyệt chưa đăng nhập vẫn có thể mở khung giao diện."),
    ("6", "Chỉ phát hiện chưa đăng nhập sau khi API trả 401", "Frontend",
     "Frontend đợi trang gọi API; khi backend trả 401, Axios interceptor mới xóa session và chuyển về /login.",
     "Trung bình", "Giao diện có thể xuất hiện trước khi chuyển trang; trang chưa gọi API có thể nằm lại."),
    ("7", "DashboardLayout mặc định người lạ là USER", "Frontend",
     "Khi không có token/user, layout dùng role mặc định USER và tên mặc định User thay vì coi session là không hợp lệ.",
     "Trung bình", "Người chưa đăng nhập vẫn được render layout người dùng."),
    ("8", "Guard admin dựa vào localStorage.role", "Frontend",
     "AdminOnly và ManagementOnly đọc role từ localStorage, trong khi người dùng có thể tự sửa dữ liệu này bằng DevTools.",
     "Trung bình về giao diện", "Có thể làm frontend hiển thị trang hoặc menu admin."),
    ("9", "Có thể giả role để mở giao diện admin", "Frontend",
     "Sửa localStorage.role thành ADMIN có thể vượt kiểm tra giao diện, nhưng không thay đổi quyền thật ở backend.",
     "Thấp - trung bình", "Lộ cấu trúc chức năng quản trị; API vẫn bị chặn."),
    ("10", "API nghiệp vụ vẫn yêu cầu JWT", "Backend - không phải lỗi",
     "Spring Security dùng anyRequest().authenticated(); document, upload, summary, quiz và các API còn lại cần JWT.",
     "An toàn", "Người lạ có thể thấy UI nhưng không đọc hoặc thay đổi dữ liệu."),
    ("11", "Giả role frontend không vượt được backend", "Backend - không phải lỗi",
     "Backend lấy role thật từ user trong database sau khi xác thực JWT, không tin role trong localStorage.",
     "An toàn", "User giả role chỉ mở được UI; API admin trả 401/403."),
    ("12", "Một số endpoint được mở công khai theo chủ đích", "Theo thiết kế",
     "Login, register, verify/reset password, GET /api/plans và tài liệu chia sẻ công khai được permitAll.",
     "Không phải lỗi", "Cần public theo nghiệp vụ; share token vẫn phải khó đoán và còn hiệu lực."),
]

OWNERS = [
    ("Backend", "Bổ sung chống brute force, rate limit đăng nhập, đếm số lần sai, khóa tạm, trả 429 khi vượt giới hạn và reset bộ đếm sau khi đăng nhập đúng."),
    ("Database", "Nếu khóa theo tài khoản, bổ sung failed_login_attempts, last_failed_login_at và locked_until."),
    ("Frontend", "Tạo guard đăng nhập chung cho /app/**; session không hợp lệ phải chuyển ngay về /login, không chờ API trả 401."),
    ("Frontend", "Không xem role trong localStorage là căn cứ bảo mật. Role phía client chỉ dùng điều khiển giao diện."),
    ("Backend", "Tiếp tục giữ xác thực và phân quyền tại API vì đây là lớp bảo mật quyết định."),
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent=0):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, (cell, width) in enumerate(zip(row.cells, widths_dxa)):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)


def set_run(run, size=9, bold=False, color="000000", italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_cell(cell, size=8.5, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.08
        for run in paragraph.runs:
            set_run(run, size=size, bold=bold, color=color)


def set_style_font(style, size, color="000000", bold=False, before=0, after=6, line=1.1):
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run(run, size=8.5, color="666666")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    # compact_reference_guide with named landscape-table override (9.7 in / 13968 DXA).
    set_style_font(doc.styles["Normal"], 9.5, after=5, line=1.12)
    set_style_font(doc.styles["Title"], 22, NAVY, True, after=5)
    set_style_font(doc.styles["Subtitle"], 10.5, "666666", False, after=12)
    set_style_font(doc.styles["Heading 1"], 15, BLUE, True, before=12, after=7)
    set_style_font(doc.styles["Heading 2"], 12, NAVY, True, before=9, after=5)

    header = section.header.paragraphs[0]
    header.text = "AI STUDY HUB  |  PHÂN TÍCH XÁC THỰC VÀ PHÂN QUYỀN"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        set_run(run, size=8.5, bold=True, color="7A7A7A")

    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.add_run("BẢNG TÓM TẮT PHÂN TÍCH LỖI BẢO MẬT")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Phạm vi: đăng nhập sai nhiều lần và truy cập trực tiếp đường dẫn chức năng khi chưa đăng nhập")

    lead = doc.add_paragraph()
    lead.paragraph_format.space_after = Pt(10)
    r = lead.add_run("Kết luận chính: ")
    set_run(r, size=10, bold=True, color=NAVY)
    r = lead.add_run(
        "Hệ thống chưa có cơ chế chống brute force và frontend thiếu guard đăng nhập chung cho /app/**. "
        "Tuy nhiên, phần lớn API nghiệp vụ vẫn được backend bảo vệ bằng JWT."
    )
    set_run(r, size=10)

    doc.add_paragraph("1. Danh sách phát hiện", style="Heading 1")
    headers = ["STT", "Lỗi / phát hiện", "Phía phụ trách", "Chi tiết", "Mức độ", "Hậu quả"]
    widths = [480, 2450, 1500, 5200, 1550, 2788]  # total 13968 DXA
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for idx, text in enumerate(headers):
        table.rows[0].cells[idx].text = text
        set_cell_shading(table.rows[0].cells[idx], NAVY)
        style_cell(table.rows[0].cells[idx], size=8.5, bold=True, color=WHITE,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])

    for item in FINDINGS:
        cells = table.add_row().cells
        for idx, text in enumerate(item):
            cells[idx].text = text
            align = WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 2, 4) else WD_ALIGN_PARAGRAPH.LEFT
            style_cell(cells[idx], size=8.2, align=align)
        if int(item[0]) % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, LIGHT_GRAY)
        severity = item[4].lower()
        color = GREEN if ("an toàn" in severity or "không phải lỗi" in severity) else RED if "cao" in severity else AMBER
        for run in cells[4].paragraphs[0].runs:
            set_run(run, size=8.2, bold=True, color=color)

    set_table_geometry(table, widths, indent=0)

    doc.add_paragraph("2. Phân chia nhiệm vụ xử lý", style="Heading 1")
    owner_table = doc.add_table(rows=1, cols=2)
    owner_table.style = "Table Grid"
    owner_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    owner_widths = [2200, 11768]
    for idx, text in enumerate(("Người/phần phụ trách", "Nhiệm vụ cần xử lý")):
        owner_table.rows[0].cells[idx].text = text
        set_cell_shading(owner_table.rows[0].cells[idx], BLUE)
        style_cell(owner_table.rows[0].cells[idx], size=9, bold=True, color=WHITE,
                   align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(owner_table.rows[0])
    for idx, (owner, task) in enumerate(OWNERS):
        cells = owner_table.add_row().cells
        cells[0].text = owner
        cells[1].text = task
        style_cell(cells[0], size=9, bold=True, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER)
        style_cell(cells[1], size=9)
        if idx % 2 == 1:
            set_cell_shading(cells[0], LIGHT_BLUE)
            set_cell_shading(cells[1], LIGHT_BLUE)
    set_table_geometry(owner_table, owner_widths, indent=0)

    doc.add_paragraph("3. Cách hiểu đúng về vấn đề truy cập URL", style="Heading 1")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    for label, text in [
        ("Mở được giao diện React: ", "có thể xảy ra vì frontend chưa chặn route /app/** ngay từ đầu."),
        ("Truy cập được dữ liệu: ", "không, nếu API tương ứng được Spring Security bảo vệ và request không có JWT hợp lệ."),
        ("Giả role trong localStorage: ", "có thể làm thay đổi giao diện nhưng không làm thay đổi quyền thật tại backend."),
    ]:
        r = p.add_run(label)
        set_run(r, size=9.5, bold=True, color=NAVY)
        r = p.add_run(text + "\n")
        set_run(r, size=9.5)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(6)
    note.paragraph_format.space_after = Pt(0)
    r = note.add_run("Ghi chú kiểm thử: ")
    set_run(r, size=8.5, bold=True, color="666666")
    r = note.add_run(
        "Kết luận được xác định từ AuthService, User entity, SecurityConfig, JwtAuthenticationFilter, "
        "router frontend và Axios interceptor. Tại thời điểm kiểm tra không có tiến trình frontend/backend "
        "lắng nghe ở cổng 5173/8080 nên chưa thực hiện kiểm thử request runtime."
    )
    set_run(r, size=8.5, italic=True, color="666666")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
